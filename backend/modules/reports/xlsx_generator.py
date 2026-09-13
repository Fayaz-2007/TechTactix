"""
Generates a simple Excel (.xlsx) report for a query/answer/sources result,
entirely in memory (no disk writes).
"""

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

EXCERPT_LENGTH = 200

_HEADER_FONT = Font(bold=True)
_SECTION_FONT = Font(bold=True, size=12)
_WRAP_ALIGNMENT = Alignment(wrap_text=True, vertical="top")


def _truncate(text: str, length: int = EXCERPT_LENGTH) -> str:
    text = (text or "").strip()
    if len(text) <= length:
        return text
    return text[:length].rstrip() + "..."


def generate_xlsx(query: str, answer: str, sources: list) -> bytes:
    """
    Build a single-sheet workbook: row 1 = Query/Answer headers, row 2 =
    their values, then a small Sources table below.

    Args:
        query: the user's original question.
        answer: the generated answer text.
        sources: list of dicts, each expected to have "document_id",
            "chunk_text" (or "text"), and "score".

    Returns:
        The raw .xlsx file bytes.
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Report"

    # --- Query / Answer block ---
    sheet["A1"] = "Query"
    sheet["B1"] = "Answer"
    sheet["A1"].font = _HEADER_FONT
    sheet["B1"].font = _HEADER_FONT

    sheet["A2"] = query or ""
    sheet["B2"] = answer or ""
    sheet["A2"].alignment = _WRAP_ALIGNMENT
    sheet["B2"].alignment = _WRAP_ALIGNMENT

    # --- Sources table ---
    section_row = 4
    sheet.cell(row=section_row, column=1, value="Sources").font = _SECTION_FONT

    header_row = section_row + 1
    headers = ["Document ID", "Chunk Excerpt", "Relevance Score"]
    for col_index, header in enumerate(headers, start=1):
        sheet.cell(row=header_row, column=col_index, value=header).font = _HEADER_FONT

    for row_offset, source in enumerate(sources, start=1):
        row = header_row + row_offset
        document_id = source.get("document_id", "unknown")
        chunk_text = source.get("chunk_text") or source.get("text", "")
        score = source.get("score", "")
        if isinstance(score, (int, float)):
            score = round(score, 4)

        sheet.cell(row=row, column=1, value=document_id)
        excerpt_cell = sheet.cell(row=row, column=2, value=_truncate(chunk_text))
        excerpt_cell.alignment = _WRAP_ALIGNMENT
        sheet.cell(row=row, column=3, value=score)

    for column_letter, width in {"A": 24, "B": 60, "C": 18}.items():
        sheet.column_dimensions[column_letter].width = width

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
