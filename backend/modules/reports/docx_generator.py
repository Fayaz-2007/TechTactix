"""
Generates a clean, simple Word (.docx) report for a query/answer/sources
result, entirely in memory (no disk writes).
"""

import io
from datetime import datetime

from docx import Document
from docx.shared import Inches, Pt, RGBColor

EXCERPT_LENGTH = 300


def _truncate(text: str, length: int = EXCERPT_LENGTH) -> str:
    text = (text or "").strip()
    if len(text) <= length:
        return text
    return text[:length].rstrip() + "..."


def generate_docx(query: str, answer: str, sources: list) -> bytes:
    """
    Build a Word document with a title, the question, the answer, and a
    Sources section (document_id + a short excerpt per source).

    Args:
        query: the user's original question.
        answer: the generated answer text.
        sources: list of dicts, each expected to have at least
            "document_id" and "chunk_text" (or "text").

    Returns:
        The raw .docx file bytes.
    """
    document = Document()

    document.add_heading("Document Assistant Report", level=0)

    meta_paragraph = document.add_paragraph()
    meta_run = meta_paragraph.add_run(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    meta_run.italic = True
    meta_run.font.size = Pt(9)
    meta_run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    document.add_heading("Question", level=1)
    document.add_paragraph(query or "(no query provided)")

    document.add_heading("Answer", level=1)
    document.add_paragraph(answer or "(no answer generated)")

    document.add_heading("Sources", level=1)
    if sources:
        for i, source in enumerate(sources, start=1):
            document_id = source.get("document_id", "unknown")
            chunk_text = source.get("chunk_text") or source.get("text", "")
            excerpt = _truncate(chunk_text)

            header_paragraph = document.add_paragraph()
            header_run = header_paragraph.add_run(f"{i}. Document: {document_id}")
            header_run.bold = True

            if excerpt:
                excerpt_paragraph = document.add_paragraph(excerpt)
                excerpt_paragraph.paragraph_format.left_indent = Inches(0.25)
    else:
        document.add_paragraph("No sources were used for this answer.")

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
