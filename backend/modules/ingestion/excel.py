"""
Text extraction for .xlsx spreadsheets, used by the ingestion dispatch in
backend/modules/ingestion/ocr.py.

Fully offline — openpyxl reads the file directly, no network calls.
"""

from openpyxl import load_workbook


def extract_excel_text(filepath: str) -> str:
    """
    Read every sheet in an .xlsx workbook and flatten it to plain text.

    Each row's cell values are joined with " | "; each sheet's rows are
    prefixed with a "Sheet: <sheet name>" line; sheets are joined with a
    blank line between them.
    """
    workbook = load_workbook(filepath, data_only=True, read_only=True)

    try:
        sheet_blocks = []
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            lines = [f"Sheet: {sheet_name}"]
            for row in sheet.iter_rows(values_only=True):
                cells = ["" if value is None else str(value) for value in row]
                lines.append(" | ".join(cells))
            sheet_blocks.append("\n".join(lines))
    finally:
        workbook.close()

    return "\n\n".join(sheet_blocks)
